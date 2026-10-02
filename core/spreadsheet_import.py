from datetime import date as date_type, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.core.exceptions import ValidationError
from django.utils import timezone
from openpyxl import load_workbook

from .models import Broker, Charge, Payment, Person, PoliceReportExport, Room, Stay, Tenancy
from .services import allocate_unallocated_payments, create_commission_charge, set_planned_checkout


class TemplateImportError(Exception):
    pass


ROOM_STATUS = {
    "空房": Room.Status.VACANT,
    "在住": Room.Status.OCCUPIED,
    "合同即将到期": Room.Status.EXPIRING,
    "即将到期": Room.Status.EXPIRING,
    "自用": Room.Status.SELF_USE,
    "维修中": Room.Status.MAINTENANCE,
}

PAYMENT_CYCLE = {
    "月付": Tenancy.PaymentCycle.MONTHLY,
    "季付": Tenancy.PaymentCycle.QUARTERLY,
    "半年付": Tenancy.PaymentCycle.HALF_YEAR,
    "年付": Tenancy.PaymentCycle.YEARLY,
}

TENANCY_STATUS = {
    "在租": Tenancy.Status.ACTIVE,
    "未入住": Tenancy.Status.UPCOMING,
    "已退租": Tenancy.Status.ENDED,
}

STAY_TYPE = {
    "常住": Stay.Type.PERMANENT,
    "探望": Stay.Type.VISITOR,
    "暂住": Stay.Type.VISITOR,
    "探望/暂住": Stay.Type.VISITOR,
    "管理员": Stay.Type.MANAGER,
}

DIRECTION = {
    "应收": Charge.Direction.INCOME,
    "应付": Charge.Direction.EXPENSE,
}

PAYMENT_DIRECTION = {
    "收": Payment.Direction.RECEIVE,
    "收入": Payment.Direction.RECEIVE,
    "应收": Payment.Direction.RECEIVE,
    "付": Payment.Direction.PAY,
    "支出": Payment.Direction.PAY,
    "应付": Payment.Direction.PAY,
}

CATEGORY = {
    "房租": Charge.Category.RENT,
    "押金": Charge.Category.DEPOSIT,
    "取暖费": Charge.Category.HEATING,
    "中介佣金": Charge.Category.COMMISSION,
    "工资": Charge.Category.WAGE,
    "产权方租金": Charge.Category.PROPERTY_RENT,
    "维修": Charge.Category.REPAIR,
    "水电等费用": Charge.Category.UTILITY,
    "换房补差": Charge.Category.MOVE_DIFF,
    "押金退款": Charge.Category.DEPOSIT_REFUND,
    "其他": Charge.Category.OTHER,
}

MONEY_DIRECTION = {
    "应收": Charge.Direction.INCOME,
    "应付": Charge.Direction.EXPENSE,
    "收": Payment.Direction.RECEIVE,
    "付": Payment.Direction.PAY,
}

FLOOR = {
    "1层": Room.Floor.FIRST,
    "一层": Room.Floor.FIRST,
    "1楼": Room.Floor.FIRST,
    "一楼": Room.Floor.FIRST,
    "1": Room.Floor.FIRST,
    "2层": Room.Floor.SECOND,
    "二层": Room.Floor.SECOND,
    "2楼": Room.Floor.SECOND,
    "二楼": Room.Floor.SECOND,
    "2": Room.Floor.SECOND,
}


def text(value):
    if value is None:
        return ""
    return str(value).strip()


def money(value, default=None):
    if value in [None, ""]:
        return default
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        raise TemplateImportError(f"金额格式不正确：{value}")


def validate_id_number(value, sheet, row_number):
    value = text(value)
    if len(value) not in {15, 18}:
        raise TemplateImportError(f"{sheet} 第 {row_number} 行身份证号需要是 15 位或 18 位：{value}")
    return value.upper()


def validate_phone(value, sheet, row_number, required=False):
    value = text(value)
    if required and not value:
        raise TemplateImportError(f"{sheet} 第 {row_number} 行手机号不能为空")
    if value and (not value.isdigit() or len(value) != 11):
        raise TemplateImportError(f"{sheet} 第 {row_number} 行电话号码需要是 11 位数字：{value}")
    return value


def floor_value(value):
    value = text(value)
    return FLOOR.get(value, Room.Floor.SECOND)


def date_value(value, required=False):
    if value in [None, ""]:
        if required:
            raise TemplateImportError("日期必填")
        return None
    if isinstance(value, date_type) and not isinstance(value, datetime):
        return value
    if hasattr(value, "date"):
        return value.date()
    if isinstance(value, str):
        value = value.strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                pass
    raise TemplateImportError(f"日期格式不正确：{value}")


def bool_value(value, default=True):
    value = text(value)
    if not value:
        return default
    return value in {"是", "Y", "y", "yes", "YES", "1", "true", "True"}


def iter_rows(workbook, sheet_name):
    if sheet_name not in workbook.sheetnames:
        return
    sheet = workbook[sheet_name]
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(cell not in [None, ""] for cell in row):
            continue
        yield row_number, row


def require(value, label, sheet, row_number):
    value = text(value)
    if not value:
        raise TemplateImportError(f"{sheet} 第 {row_number} 行缺少必填字段：{label}")
    return value


def map_choice(value, choices, label, sheet, row_number, default=None):
    value = text(value)
    if not value and default is not None:
        return default
    if value not in choices:
        raise TemplateImportError(f"{sheet} 第 {row_number} 行 {label} 不合法：{value}")
    return choices[value]


def clear_business_data():
    PoliceReportExport.objects.all().delete()
    Payment.objects.all().delete()
    Charge.objects.all().delete()
    Stay.objects.all().delete()
    Tenancy.objects.all().delete()
    Broker.objects.all().delete()
    Person.objects.all().delete()
    Room.objects.all().delete()


def person_from_fields(parts, sheet, row_number):
    parts = [text(part) for part in parts]
    while len(parts) < 7:
        parts.append("")
    name, id_number, phone, emergency_name, emergency_phone, emergency_address, notes = parts[:7]
    if not name and not id_number:
        return None
    if not name or not id_number:
        raise TemplateImportError(f"{sheet} 第 {row_number} 行人员格式需要至少包含 姓名|身份证号")
    id_number = validate_id_number(id_number, sheet, row_number)
    phone = validate_phone(phone, sheet, row_number)
    emergency_phone = validate_phone(emergency_phone, sheet, row_number)
    person, _ = Person.objects.update_or_create(
        id_number=id_number,
        defaults={
            "name": name,
            "phone": phone,
            "emergency_name": emergency_name,
            "emergency_phone": emergency_phone,
            "emergency_address": emergency_address,
            "notes": notes,
        },
    )
    return person


def parse_person_list(value, sheet, row_number):
    people = []
    raw = text(value)
    if not raw:
        return people
    for entry in raw.replace("；", ";").split(";"):
        entry = entry.strip()
        if not entry:
            continue
        people.append(person_from_fields(entry.split("|"), sheet, row_number))
    return [person for person in people if person]


def upsert_stay(person, room, stay_type, start_date, end_date=None, tenancy=None, report_note="", notes="", is_active=True):
    stay, _ = Stay.objects.update_or_create(
        person=person,
        room=room,
        stay_type=stay_type,
        defaults={
            "tenancy": tenancy if stay_type == Stay.Type.PERMANENT else None,
            "start_date": start_date,
            "end_date": end_date,
            "is_active": is_active,
            "report_note": report_note,
            "notes": notes,
        },
    )
    return stay


def payment_category_from_charge_category(category):
    if category in {Charge.Category.DEPOSIT, Charge.Category.DEPOSIT_REFUND}:
        return Payment.Category.DEPOSIT
    if category == Charge.Category.RENT:
        return Payment.Category.RENT
    return Payment.Category.OTHER


def import_split_room_template(workbook, clear=False):
    if clear:
        clear_business_data()
    counts = {"房间": 0, "人员": 0, "合同": 0, "入住": 0, "账单": 0, "收付款": 0}
    today = timezone.localdate()
    tenancies_by_room = {}

    for row_number, row in iter_rows(workbook, "房间与合同") or []:
        row = list(row) + [""] * 32
        room_number = require(row[0], "房间号", "房间与合同", row_number)
        room_defaults = {
            "status": map_choice(row[1], ROOM_STATUS, "房态", "房间与合同", row_number, default=Room.Status.VACANT),
            "listing_price": money(row[2]),
            "commission_base": money(row[3], Decimal("2000.00")),
            "orientation": text(row[4]),
            "floor": floor_value(row[5]),
            "area": money(row[6]),
            "room_password": text(row[7]),
            "water_fee": text(row[8]) or "9.5/吨",
            "electricity_fee": text(row[9]) or "1.4/度",
            "property_fee": text(row[10]) or "免",
            "internet_fee": text(row[11]) or "免",
            "heating_fee": money(row[12], Decimal("380.00")),
            "parking_fee": money(row[13], Decimal("150.00")),
            "notes": text(row[14]),
        }
        room, _ = Room.objects.update_or_create(number=room_number, defaults=room_defaults)
        counts["房间"] += 1

        main_person = person_from_fields(
            [row[15], row[16], row[17], row[18], row[19], row[20], ""],
            "房间与合同",
            row_number,
        )
        if main_person:
            counts["人员"] += 1

        contract_start = date_value(row[21])
        contract_end = date_value(row[22])
        billing_start = date_value(row[23])
        monthly_rent = money(row[24])
        if main_person and contract_start and contract_end and monthly_rent is not None:
            broker = None
            if text(row[27]):
                broker, _ = Broker.objects.get_or_create(name=text(row[27]))
            tenancy, _ = Tenancy.objects.update_or_create(
                room=room,
                primary_person=main_person,
                start_date=contract_start,
                defaults={
                    "end_date": contract_end,
                    "billing_start_date": billing_start or contract_start,
                    "monthly_rent": monthly_rent,
                    "payment_cycle": map_choice(row[25], PAYMENT_CYCLE, "付款周期", "房间与合同", row_number, default=Tenancy.PaymentCycle.MONTHLY),
                    "deposit_amount": money(row[26], Decimal("3500.00")),
                    "broker": broker,
                    "commission_base": room.commission_base,
                    "commission_manual_amount": money(row[28]),
                    "police_report_end_date": date_value(row[29]),
                    "status": map_choice(row[30], TENANCY_STATUS, "合同状态", "房间与合同", row_number, default=Tenancy.Status.ACTIVE),
                    "notes": text(row[31]),
                },
            )
            tenancies_by_room[room.number] = tenancy
            counts["合同"] += 1
            if broker:
                create_commission_charge(tenancy)
            upsert_stay(
                main_person,
                room,
                Stay.Type.PERMANENT,
                contract_start,
                tenancy.end_date if tenancy.status == Tenancy.Status.ENDED else None,
                tenancy=tenancy,
                is_active=tenancy.status == Tenancy.Status.ACTIVE,
            )
            counts["入住"] += 1
            planned_move_out_date = date_value(row[32])
            if tenancy.status == Tenancy.Status.ACTIVE and planned_move_out_date:
                refund_amount = money(row[33], tenancy.deposit_amount)
                set_planned_checkout(
                    tenancy,
                    planned_date=planned_move_out_date,
                    refund_deposit_amount=refund_amount,
                    note=text(row[34]),
                )
                if refund_amount:
                    counts["账单"] += 1

    for row_number, row in iter_rows(workbook, "入住人员") or []:
        row = list(row) + [""] * 13
        room_number = require(row[0], "房间号", "入住人员", row_number)
        room = Room.objects.filter(number=room_number).first()
        if not room:
            raise TemplateImportError(f"入住人员 第 {row_number} 行找不到房间：{room_number}")
        stay_type = map_choice(row[1], STAY_TYPE, "入住类型", "入住人员", row_number)
        person = person_from_fields(
            [row[2], row[3], row[4], row[5], row[6], row[7], row[12]],
            "入住人员",
            row_number,
        )
        if not person:
            continue
        counts["人员"] += 1
        start_date = date_value(row[8])
        end_date = date_value(row[9])
        tenancy = tenancies_by_room.get(room.number) or room.active_tenancy()
        upsert_stay(
            person,
            room,
            stay_type,
            start_date,
            end_date=end_date,
            tenancy=tenancy if stay_type == Stay.Type.PERMANENT else None,
            report_note=text(row[11]) or ("管理员" if stay_type == Stay.Type.MANAGER else "探望、暂住" if stay_type == Stay.Type.VISITOR else ""),
            notes=text(row[12]),
            is_active=bool_value(row[10], default=True),
        )
        counts["入住"] += 1

    for row_number, row in iter_rows(workbook, "期初账务") or []:
        row = list(row) + [""] * 7
        amount = money(row[4])
        if not amount:
            continue
        direction = map_choice(row[0], MONEY_DIRECTION, "方向", "期初账务", row_number)
        category = map_choice(row[1], CATEGORY, "类别", "期初账务", row_number, default=Charge.Category.OTHER)
        room = Room.objects.filter(number=text(row[2])).first() if text(row[2]) else None
        if text(row[2]) and not room:
            raise TemplateImportError(f"期初账务 第 {row_number} 行找不到房间：{text(row[2])}")
        entry_date = date_value(row[3], required=True)
        tenancy = room.active_tenancy(entry_date) if room else None
        description = text(row[5]) or dict(CATEGORY).get(text(row[1]), "期初账务")
        if direction in {Payment.Direction.RECEIVE, Payment.Direction.PAY}:
            Payment.objects.create(
                direction=direction,
                category=payment_category_from_charge_category(category),
                date=entry_date,
                room=room,
                person=tenancy.primary_person if tenancy else None,
                tenancy=tenancy,
                amount=amount,
                memo=description if not text(row[6]) else f"{description}：{text(row[6])}",
            )
            counts["收付款"] += 1
        else:
            Charge.objects.create(
                direction=direction,
                category=category,
                room=room,
                person=tenancy.primary_person if tenancy else None,
                tenancy=tenancy,
                due_date=entry_date,
                amount=amount,
                description=description,
                notes=text(row[6]),
                source=Charge.Source.OPENING,
            )
            counts["账单"] += 1

    allocate_unallocated_payments()
    for room in Room.objects.all():
        room.refresh_status(save=True)
    return counts


def import_room_summary(workbook, clear=False):
    if clear:
        clear_business_data()
    counts = {"房间": 0, "人员": 0, "合同": 0, "入住": 0, "账单": 0, "收付款": 0}
    today = timezone.localdate()
    for row_number, row in iter_rows(workbook, "房间总表") or []:
        row = list(row) + [""] * 41
        room_number = require(row[0], "房间号", "房间总表", row_number)
        room_defaults = {
            "status": map_choice(row[1], ROOM_STATUS, "房态", "房间总表", row_number, default=Room.Status.VACANT),
            "listing_price": money(row[2]),
            "commission_base": money(row[3], Decimal("2000.00")),
            "orientation": text(row[4]),
            "floor": floor_value(row[5]),
            "area": money(row[6]),
            "room_password": text(row[7]),
            "water_fee": text(row[8]) or "9.5/吨",
            "electricity_fee": text(row[9]) or "1.4/度",
            "property_fee": text(row[10]) or "免",
            "internet_fee": text(row[11]) or "免",
            "heating_fee": money(row[12], Decimal("380.00")),
            "parking_fee": money(row[13], Decimal("150.00")),
            "qr_code_url": text(row[14]),
            "notes": text(row[15]),
        }
        room, _ = Room.objects.update_or_create(number=room_number, defaults=room_defaults)
        counts["房间"] += 1

        main_person = person_from_fields(
            [row[16], row[17], row[18], row[19], row[20], row[21], ""],
            "房间总表",
            row_number,
        )
        if main_person:
            counts["人员"] += 1

        contract_start = date_value(row[22])
        contract_end = date_value(row[23])
        billing_start = date_value(row[24])
        monthly_rent = money(row[25])
        tenancy = None
        if main_person and contract_start and contract_end and monthly_rent is not None:
            broker = None
            if text(row[28]):
                broker, _ = Broker.objects.get_or_create(name=text(row[28]))
            tenancy, _ = Tenancy.objects.update_or_create(
                room=room,
                primary_person=main_person,
                start_date=contract_start,
                defaults={
                    "end_date": contract_end,
                    "billing_start_date": billing_start,
                    "monthly_rent": monthly_rent,
                    "payment_cycle": map_choice(row[26], PAYMENT_CYCLE, "付款周期", "房间总表", row_number, default=Tenancy.PaymentCycle.MONTHLY),
                    "deposit_amount": money(row[27], Decimal("3500.00")),
                    "broker": broker,
                    "commission_base": room.commission_base,
                    "commission_manual_amount": money(row[29]),
                    "police_report_end_date": date_value(row[30]),
                    "status": map_choice(row[31], TENANCY_STATUS, "合同状态", "房间总表", row_number, default=Tenancy.Status.ACTIVE),
                    "notes": text(row[32]),
                },
            )
            counts["合同"] += 1
            if broker:
                create_commission_charge(tenancy)
            upsert_stay(
                main_person,
                room,
                Stay.Type.PERMANENT,
                contract_start,
                tenancy.end_date if tenancy.status == Tenancy.Status.ENDED else None,
                tenancy=tenancy,
                is_active=tenancy.status == Tenancy.Status.ACTIVE,
            )
            counts["入住"] += 1

        default_start = contract_start or today
        for person in parse_person_list(row[33], "房间总表", row_number):
            counts["人员"] += 1
            upsert_stay(person, room, Stay.Type.PERMANENT, default_start, tenancy=tenancy)
            counts["入住"] += 1
        for person in parse_person_list(row[34], "房间总表", row_number):
            counts["人员"] += 1
            upsert_stay(person, room, Stay.Type.VISITOR, None, report_note="探望、暂住")
            counts["入住"] += 1
        for person in parse_person_list(row[35], "房间总表", row_number):
            counts["人员"] += 1
            upsert_stay(person, room, Stay.Type.MANAGER, None, report_note="管理员")
            counts["入住"] += 1

        due_date = billing_start or contract_start or today
        account_note = text(row[40])
        opening_specs = [
            (row[36], Charge.Direction.INCOME, "期初应收", "期初应收", True),
            (row[37], Charge.Direction.EXPENSE, "期初应付", "期初应付", True),
            (row[38], Payment.Direction.RECEIVE, "期初已收", "期初已收", False),
            (row[39], Payment.Direction.PAY, "期初已付", "期初已付", False),
        ]
        for value, direction, description, memo, is_charge in opening_specs:
            amount = money(value)
            if not amount:
                continue
            if is_charge:
                Charge.objects.create(
                    direction=direction,
                    category=Charge.Category.OTHER,
                    room=room,
                    person=main_person,
                    tenancy=tenancy,
                    due_date=due_date,
                    amount=amount,
                    description=description,
                    notes=account_note,
                    source=Charge.Source.OPENING,
                )
                counts["账单"] += 1
            else:
                Payment.objects.create(
                    direction=direction,
                    date=due_date,
                    room=room,
                    person=main_person,
                    tenancy=tenancy,
                    amount=amount,
                    memo=memo if not account_note else f"{memo}：{account_note}",
                )
                counts["收付款"] += 1
    allocate_unallocated_payments()
    for room in Room.objects.all():
        room.refresh_status(save=True)
    return counts


@transaction.atomic
def import_template(file_obj, clear=False):
    try:
        return _import_template(file_obj, clear=clear)
    except ValidationError as exc:
        raise TemplateImportError("资料未导入：" + "；".join(exc.messages)) from exc


def _import_template(file_obj, clear=False):
    workbook = load_workbook(file_obj, data_only=True)
    if "房间与合同" in workbook.sheetnames:
        return import_split_room_template(workbook, clear=clear)
    if "房间总表" in workbook.sheetnames:
        return import_room_summary(workbook, clear=clear)
    if not set(workbook.sheetnames) & {"房间", "人员", "合同", "入住", "账单", "收付款"}:
        raise TemplateImportError("无法识别项目模板。请在“更多 → 批量导入”下载模板；原悦山四表需先按 yueshan_check 核对，不能按旧列位置直接导入。")
    if clear:
        clear_business_data()

    counts = {"房间": 0, "人员": 0, "合同": 0, "入住": 0, "账单": 0, "收付款": 0}

    for row_number, row in iter_rows(workbook, "房间") or []:
        number = require(row[0], "房间号", "房间", row_number)
        status = map_choice(row[1], ROOM_STATUS, "房态", "房间", row_number, default=Room.Status.VACANT)
        defaults = {
            "status": status,
            "listing_price": money(row[2]),
            "commission_base": money(row[3], Decimal("2000.00")),
            "orientation": text(row[4]),
            "floor": floor_value(row[5]),
            "area": money(row[6]),
            "room_password": text(row[7]),
            "water_fee": text(row[8]) or "9.5/吨",
            "electricity_fee": text(row[9]) or "1.4/度",
            "property_fee": text(row[10]) or "免",
            "internet_fee": text(row[11]) or "免",
            "heating_fee": money(row[12], Decimal("380.00")),
            "parking_fee": money(row[13], Decimal("150.00")),
            "qr_code_url": text(row[14]),
            "notes": text(row[15]),
        }
        Room.objects.update_or_create(number=number, defaults=defaults)
        counts["房间"] += 1

    for row_number, row in iter_rows(workbook, "人员") or []:
        id_number = require(row[0], "身份证号", "人员", row_number)
        name = require(row[1], "姓名", "人员", row_number)
        Person.objects.update_or_create(
            id_number=id_number,
            defaults={
                "name": name,
                "phone": text(row[2]),
                "emergency_name": text(row[3]),
                "emergency_phone": text(row[4]),
                "emergency_address": text(row[5]),
                "notes": text(row[6]),
            },
        )
        counts["人员"] += 1

    for row_number, row in iter_rows(workbook, "合同") or []:
        room_number = require(row[0], "房间号", "合同", row_number)
        id_number = require(row[1], "主租客身份证号", "合同", row_number)
        room = Room.objects.filter(number=room_number).first()
        person = Person.objects.filter(id_number=id_number).first()
        if not room:
            raise TemplateImportError(f"合同 第 {row_number} 行找不到房间：{room_number}")
        if not person:
            raise TemplateImportError(f"合同 第 {row_number} 行找不到人员身份证：{id_number}")
        broker = None
        if text(row[8]):
            broker, _ = Broker.objects.get_or_create(name=text(row[8]))
        tenancy, _ = Tenancy.objects.update_or_create(
            room=room,
            primary_person=person,
            start_date=date_value(row[2], required=True),
            defaults={
                "end_date": date_value(row[3], required=True),
                "billing_start_date": date_value(row[4]),
                "monthly_rent": money(row[5], Decimal("0.00")),
                "payment_cycle": map_choice(row[6], PAYMENT_CYCLE, "付款周期", "合同", row_number, default=Tenancy.PaymentCycle.MONTHLY),
                "deposit_amount": money(row[7], Decimal("3500.00")),
                "broker": broker,
                "commission_base": room.commission_base,
                "commission_manual_amount": money(row[9]),
                "police_report_end_date": date_value(row[10]),
                "status": map_choice(row[11], TENANCY_STATUS, "状态", "合同", row_number, default=Tenancy.Status.ACTIVE),
                "notes": text(row[12]),
            },
        )
        Stay.objects.update_or_create(
            person=person,
            room=room,
            tenancy=tenancy,
            defaults={
                "stay_type": Stay.Type.PERMANENT,
                "start_date": tenancy.start_date,
                "end_date": tenancy.end_date if tenancy.status == Tenancy.Status.ENDED else None,
                "is_active": tenancy.status == Tenancy.Status.ACTIVE,
            },
        )
        if broker:
            create_commission_charge(tenancy)
        counts["合同"] += 1

    for row_number, row in iter_rows(workbook, "入住") or []:
        room_number = require(row[0], "房间号", "入住", row_number)
        id_number = require(row[1], "身份证号", "入住", row_number)
        room = Room.objects.filter(number=room_number).first()
        person = Person.objects.filter(id_number=id_number).first()
        if not room:
            raise TemplateImportError(f"入住 第 {row_number} 行找不到房间：{room_number}")
        if not person:
            raise TemplateImportError(f"入住 第 {row_number} 行找不到人员身份证：{id_number}")
        stay_type = map_choice(row[2], STAY_TYPE, "入住类型", "入住", row_number)
        start = date_value(row[3])
        tenancy = room.active_tenancy() if stay_type == Stay.Type.PERMANENT else None
        Stay.objects.update_or_create(
            person=person,
            room=room,
            tenancy=tenancy,
            defaults={
                "stay_type": stay_type,
                "start_date": start,
                "end_date": date_value(row[4]),
                "is_active": bool_value(row[5], default=True),
                "report_note": text(row[6]),
                "notes": text(row[7]),
            },
        )
        counts["入住"] += 1

    for row_number, row in iter_rows(workbook, "账单") or []:
        room = Room.objects.filter(number=text(row[2])).first() if text(row[2]) else None
        person = Person.objects.filter(id_number=text(row[3])).first() if text(row[3]) else None
        tenancy = room.active_tenancy(date_value(row[4], required=True)) if room else None
        Charge.objects.create(
            direction=map_choice(row[0], DIRECTION, "方向", "账单", row_number),
            category=map_choice(row[1], CATEGORY, "类别", "账单", row_number),
            room=room,
            person=person or (tenancy.primary_person if tenancy else None),
            tenancy=tenancy,
            due_date=date_value(row[4], required=True),
            period_start=date_value(row[5]),
            period_end=date_value(row[6]),
            amount=money(row[7], Decimal("0.00")),
            description=require(row[8], "说明", "账单", row_number),
            notes=text(row[9]),
            source=Charge.Source.MANUAL,
        )
        counts["账单"] += 1

    for row_number, row in iter_rows(workbook, "收付款") or []:
        room_number = require(row[2], "房间号", "收付款", row_number)
        room = Room.objects.filter(number=room_number).first()
        if not room:
            raise TemplateImportError(f"收付款 第 {row_number} 行找不到房间：{room_number}")
        payment_date = date_value(row[1], required=True)
        tenancy = room.active_tenancy(payment_date)
        Payment.objects.create(
            direction=map_choice(row[0], PAYMENT_DIRECTION, "方向", "收付款", row_number),
            date=payment_date,
            room=room,
            person=tenancy.primary_person if tenancy else None,
            tenancy=tenancy,
            amount=money(row[3], Decimal("0.00")),
            memo=text(row[4]),
        )
        counts["收付款"] += 1

    allocate_unallocated_payments()
    for room in Room.objects.all():
        room.refresh_status(save=True)
    return counts
