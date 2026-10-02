from datetime import date
from importlib import import_module
from io import BytesIO
from unittest.mock import patch

from django.apps import apps
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .exports import police_report_rows, police_report_workbook
from .models import Charge, Payment, Person, PoliceReportExport, Room, Stay, Tenancy
from .services import checkout_tenancy, renew_tenancy, set_planned_checkout
from .spreadsheet_import import clear_business_data
from .views import police_fingerprint
from .test_support import authenticated_client, configure_test_account


class PoliceReportingTests(TestCase):
    def setUp(self):
        configure_test_account(self)
        self.clock = patch("django.utils.timezone.localdate", return_value=date(2026, 10, 2))
        self.clock.start()
        self.addCleanup(self.clock.stop)
        self.room = Room.objects.create(number="A01", listing_price=3000)
        self.person = Person.objects.create(name="测试主租客", id_number="110101199001010011", phone="013800000000")
        self.tenancy = Tenancy.objects.create(room=self.room, primary_person=self.person,
            start_date=date(2026, 6, 1), end_date=date(2026, 8, 31), monthly_rent=3000,
            billing_enabled=False, deposit_amount=None)
        self.stay = Stay.objects.create(person=self.person, room=self.room, tenancy=self.tenancy,
            start_date=date(2026, 6, 4), report_note="2026.6.4-2026.8.31租约延长")

    def new_stay(self, name="测试暂住人", **kwargs):
        person = Person.objects.create(name=name)
        return Stay.objects.create(person=person, room=self.room, **kwargs)

    def depart(self, stay=None):
        stay = stay or self.stay
        stay.is_active = False
        stay.end_date = date(2026, 10, 1)
        stay.save(update_fields=["is_active", "end_date"])
        return stay

    def generate(self, client=None):
        client = client or self.client
        response = client.get(reverse("export_police"))
        self.assertEqual(response.status_code, 200)
        return PoliceReportExport.objects.first(), response

    def confirm(self, report, client=None):
        response = (client or self.client).post(reverse("police_report_confirm", args=[report.pk]))
        self.assertEqual(response.status_code, 302)

    def preview_payload(self, changes=None, action="save"):
        preview = self.client.get(reverse("police_preview"))
        formset = preview.context["formset"]
        payload = {"action": action}
        for name, value in formset.management_form.initial.items():
            payload[f"form-{name}"] = str(value)
        for form in formset:
            values = dict(form.initial)
            values.update((changes or {}).get(values["stay_id"], {}))
            payload.update({form.add_prefix(name): str(value) for name, value in values.items()})
        return payload

    def test_homepage_opens_editable_report_preview(self):
        response = self.client.get(reverse("dashboard"))
        self.assertContains(response, f'href="{reverse("police_preview")}">生成报备表</a>')
        preview = self.client.get(reverse("police_preview"))
        self.assertContains(preview, "保存调整")
        self.assertContains(preview, "保存并生成报备表")

    def test_manual_final_text_is_saved_and_exported_without_changing_source_records(self):
        text = "2026.06.04-2026.12.03 人工核对"
        payload = self.preview_payload({self.stay.pk: {"text": text}}, action="generate")
        response = self.client.post(reverse("police_preview"), payload)
        report = PoliceReportExport.objects.get()
        self.assertRedirects(response, reverse("police_report_detail", args=[report.pk]) + "?download=1")
        self.assertEqual(report.rows[0]["values"][6], text)
        downloaded = authenticated_client(self).get(reverse("police_report_download", args=[report.pk]))
        self.assertEqual(load_workbook(BytesIO(downloaded.content)).active["G3"].value, text)
        self.stay.refresh_from_db(); self.tenancy.refresh_from_db()
        self.assertEqual(self.stay.police_report_text_override, text)
        self.assertEqual(self.stay.report_note, "2026.6.4-2026.8.31租约延长")
        self.assertEqual(self.tenancy.end_date, date(2026, 8, 31))
        self.assertFalse(Charge.objects.exists()); self.assertFalse(Payment.objects.exists())

    def test_unchanged_preview_does_not_freeze_automatic_content_and_blank_restores_it(self):
        self.client.post(reverse("police_preview"), self.preview_payload())
        self.stay.refresh_from_db()
        self.assertEqual(self.stay.police_report_text_override, "")
        self.client.post(reverse("police_preview"), self.preview_payload({self.stay.pk: {"text": "人工说明"}}))
        self.assertEqual(police_report_rows()[0]["text"], "人工说明")
        self.client.post(reverse("police_preview"), self.preview_payload({self.stay.pk: {"text": ""}}))
        self.assertEqual(police_report_rows()[0]["text"], "2026.06.01-2027.05.31租约延长")

    def test_excluded_departure_persists_across_clients_and_can_be_selected_again(self):
        self.depart()
        payload = self.preview_payload({self.stay.pk: {"report_departure": "no"}})
        self.assertRedirects(self.client.post(reverse("police_preview"), payload), reverse("police_preview"))
        self.stay.refresh_from_db()
        self.assertFalse(self.stay.police_departure_required)
        self.assertIsNone(self.stay.police_departure_reported_at)
        self.assertEqual(police_report_rows(), [])
        preview = authenticated_client(self).get(reverse("police_preview"))
        self.assertEqual(preview.context["excluded_count"], 1)
        self.assertContains(preview, "已保存：不报备")
        report, _ = self.generate(authenticated_client(self)); self.confirm(report)
        self.stay.refresh_from_db()
        self.assertIsNone(self.stay.police_departure_reported_at)
        self.client.post(reverse("police_preview"), self.preview_payload({self.stay.pk: {"report_departure": "yes"}}))
        self.assertEqual(police_report_rows()[0]["text"], "退租")

    def test_generate_applies_unsaved_choices_and_keeps_old_snapshot(self):
        self.depart()
        current = self.new_stay("当前暂住", stay_type="visitor")
        old_report, _ = self.generate()
        payload = self.preview_payload({self.stay.pk: {"report_departure": "no"},
                                        current.pk: {"text": "假期探亲"}}, action="generate")
        self.client.post(reverse("police_preview"), payload)
        report = PoliceReportExport.objects.first()
        self.assertEqual(len(report.rows), 1)
        self.assertEqual(report.rows[0]["values"][6], "假期探亲")
        old_report.refresh_from_db()
        self.assertEqual(old_report.departure_count, 1)
        self.assertEqual(len(old_report.rows), 2)

    def test_invalid_edit_saves_neither_other_rows_nor_departure_choices(self):
        self.depart()
        current = self.new_stay("当前暂住", stay_type="visitor")
        payload = self.preview_payload({self.stay.pk: {"report_departure": "no"},
                                        current.pk: {"text": "2026.02.30-2026.03.01"}}, action="generate")
        response = self.client.post(reverse("police_preview"), payload)
        self.assertContains(response, "报备内容包含不存在的日期")
        self.stay.refresh_from_db(); current.refresh_from_db()
        self.assertTrue(self.stay.police_departure_required)
        self.assertEqual(current.police_report_text_override, "")
        self.assertFalse(PoliceReportExport.objects.exists())

    def test_stale_departure_edit_does_not_change_new_event(self):
        self.depart()
        payload = self.preview_payload({self.stay.pk: {"report_departure": "no", "text": "旧退租说明"}})
        self.stay.end_date = date(2026, 10, 2); self.stay.save()
        response = self.client.post(reverse("police_preview"), payload)
        self.assertContains(response, "人员或退租状态已变化")
        self.stay.refresh_from_db()
        self.assertTrue(self.stay.police_departure_required)
        self.assertEqual(self.stay.police_report_text_override, "")

    def test_incomplete_or_tampered_edit_cannot_generate_or_save(self):
        self.depart()
        payload = self.preview_payload({self.stay.pk: {"report_departure": "no"}})
        for change in ({"form-TOTAL_FORMS": "0"}, {"form-0-stay_id": "9999"}, {"form-0-report_departure": ""}):
            with self.subTest(change=change):
                response = self.client.post(reverse("police_preview"), {**payload, **change})
                self.assertEqual(response.status_code, 200)
                self.stay.refresh_from_db(); self.assertTrue(self.stay.police_departure_required)
                self.assertFalse(PoliceReportExport.objects.exists())
        self.assertEqual(self.client.post(reverse("police_preview"), {"action": "generate"}).status_code, 200)
        self.assertFalse(PoliceReportExport.objects.exists())

    def test_new_departure_defaults_to_reporting_and_discards_prior_manual_text(self):
        self.depart()
        self.client.post(reverse("police_preview"), self.preview_payload({self.stay.pk: {
            "report_departure": "no", "text": "旧事件说明"}}))
        self.stay.refresh_from_db()
        self.stay.is_active = True; self.stay.end_date = None; self.stay.save()
        self.assertEqual(self.stay.police_report_text_override, "")
        self.client.post(reverse("police_preview"), self.preview_payload({self.stay.pk: {"text": "在住说明"}}))
        checkout_tenancy(self.tenancy, checkout_date=date(2026, 10, 2))
        self.stay.refresh_from_db()
        self.assertTrue(self.stay.police_departure_required)
        self.assertEqual(self.stay.police_report_text_override, "")
        self.assertEqual(police_report_rows()[0]["text"], "退租")

    def test_short_contract_uses_contract_start_and_exact_year_without_writing_sources(self):
        self.tenancy.police_report_start_date = date(2026, 6, 4)
        self.tenancy.police_report_end_date = date(2028, 6, 3)
        self.tenancy.save()
        expected = "2026.06.01-2027.05.31租约延长"
        self.assertEqual(police_report_rows()[0]["text"], expected)
        self.tenancy.refresh_from_db(); self.stay.refresh_from_db()
        self.assertEqual(self.tenancy.end_date, date(2026, 8, 31))
        self.assertEqual(self.tenancy.police_report_start_date, date(2026, 6, 4))
        self.assertEqual(self.tenancy.police_report_end_date, date(2028, 6, 3))
        self.assertEqual(self.stay.report_note, "2026.6.4-2026.8.31租约延长")
        self.assertFalse(Charge.objects.exists()); self.assertFalse(Payment.objects.exists())

    def test_short_roommate_and_unlinked_roommate_use_same_contract(self):
        linked = self.new_stay("同住人", tenancy=self.tenancy, report_note="2026.7.1-2026.8.31同住")
        unlinked = self.new_stay("无关联合同同住人", report_note="同住", start_date=None)
        rows = {r["stay"].pk: r["text"] for r in police_report_rows()}
        self.assertEqual(rows[linked.pk], "2026.06.01-2027.05.31同住")
        self.assertEqual(rows[unlinked.pk], "2026.06.01-2027.05.31 同住")

    def test_six_month_contract_is_not_extended_even_if_person_just_arrived(self):
        self.tenancy.end_date = date(2026, 11, 30)
        self.tenancy.save()
        self.stay.report_note = ""; self.stay.start_date = date(2026, 10, 1); self.stay.save()
        self.assertEqual(police_report_rows()[0]["text"], "2026.06.01-2026.11.30")

    def test_all_current_types_and_unknown_or_expired_visitor_dates_are_included(self):
        holiday = self.new_stay("假期暂住人", stay_type="visitor", start_date=None, report_note="假期暂住")
        visitor = self.new_stay("到期暂住人", tenancy=self.tenancy, stay_type="visitor",
            start_date=date(2026, 7, 1), end_date=date(2026, 7, 15), report_note="2026.7.1-2026.7.15探亲")
        manager = self.new_stay("管理员", stay_type="manager", start_date=None)
        future = self.new_stay("未来住户", stay_type="visitor", start_date=date(2026, 11, 1))
        rows = {r["stay"].pk: r["text"] for r in police_report_rows()}
        self.assertEqual(set(rows), {self.stay.pk, holiday.pk, visitor.pk, manager.pk})
        self.assertEqual(rows[visitor.pk], "2026.7.1-2026.7.15探亲")
        self.assertEqual(rows[holiday.pk], "假期暂住")
        self.assertNotIn(future.pk, rows)

    def test_sample_seven_columns_and_identifier_text_survive_round_trip(self):
        report, response = self.generate()
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual({str(r) for r in sheet.merged_cells.ranges}, {"A1:G1"})
        self.assertEqual([c.value for c in sheet[2]], ["姓名", "身份证", "房间号", "电话", "紧急联系人", "紧急联系人电话", "合同时间"])
        self.assertEqual(sheet.max_column, 7)
        self.assertEqual(sheet["B3"].data_type, "s")
        self.assertEqual(sheet["D3"].value, "013800000000")
        self.assertEqual(report.rows[0]["values"][6], "2026.06.01-2027.05.31租约延长")

    def test_literal_text_cannot_become_excel_formula(self):
        self.person.name = "=1+1"; self.person.save()
        _, response = self.generate()
        cell = load_workbook(BytesIO(response.content)).active["A3"]
        self.assertEqual(cell.value, "=1+1"); self.assertEqual(cell.data_type, "s")

    def test_preview_and_multiple_downloads_keep_departure_pending_until_sent(self):
        self.depart()
        for _ in range(2):
            preview = self.client.get(reverse("police_preview"))
            self.assertEqual(preview.context["departure_count"], 1)
        self.assertFalse(PoliceReportExport.objects.exists())
        report, _ = self.generate()
        self.generate()
        self.assertEqual(police_report_rows()[0]["text"], "退租")
        self.stay.refresh_from_db(); self.assertIsNone(self.stay.police_departure_reported_at)
        self.confirm(report, authenticated_client(self))
        self.assertEqual(police_report_rows(), [])
        self.stay.refresh_from_db(); self.assertIsNotNone(self.stay.police_departure_reported_at)
        self.confirm(report)
        self.assertEqual(police_report_rows(), [])

    def test_confirming_older_report_does_not_consume_new_departure(self):
        report, _ = self.generate()
        self.depart()
        self.confirm(report)
        self.assertEqual(police_report_rows()[0]["text"], "退租")

    def test_confirming_batch_only_consumes_its_own_departure_people(self):
        self.depart()
        report, _ = self.generate()
        later = self.new_stay("后来离开的人", stay_type="visitor", start_date=None)
        self.depart(later)
        self.confirm(report)
        self.assertEqual([r["stay"].pk for r in police_report_rows()], [later.pk])

    def test_reactivation_then_departure_needs_new_token_and_report(self):
        self.depart(); report, _ = self.generate()
        old_token = self.stay.police_departure_token
        self.stay.is_active = True; self.stay.save()
        self.depart()
        self.assertNotEqual(self.stay.police_departure_token, old_token)
        self.confirm(report)
        self.assertEqual(police_report_rows()[0]["text"], "退租")

    def test_departure_date_correction_is_not_consumed_by_stale_batch(self):
        self.depart(); report, _ = self.generate()
        self.stay.end_date = date(2026, 9, 30); self.stay.save()
        self.confirm(report)
        self.assertEqual(police_report_rows()[0]["text"], "退租")

    def test_repeated_stale_object_save_does_not_erase_confirmation(self):
        self.depart(); report, _ = self.generate()
        self.confirm(report)
        self.stay.notes = "更新备注"; self.stay.save()
        self.assertEqual(police_report_rows(), [])

    def test_replaying_remove_after_confirmation_cannot_create_a_fake_departure(self):
        self.depart(); report, _ = self.generate(); self.confirm(report)
        self.stay.refresh_from_db(); old_token = self.stay.police_departure_token
        with patch("django.utils.timezone.localdate", return_value=date(2026, 10, 3)):
            for route, args in [("person_stay_remove", [self.person.pk, self.stay.pk]),
                                ("room_person_remove", [self.room.pk, self.stay.pk])]:
                self.assertEqual(self.client.post(reverse(route, args=args)).status_code, 302)
        self.stay.refresh_from_db()
        self.assertEqual(self.stay.police_departure_token, old_token)
        self.assertEqual(self.stay.end_date, date(2026, 10, 1))
        self.assertEqual(police_report_rows(), [])

    def test_same_person_can_report_old_departure_and_new_current_stay(self):
        self.depart()
        other = Room.objects.create(number="A02", listing_price=3000)
        Stay.objects.create(person=self.person, room=other, start_date=date(2026, 10, 2))
        rows = police_report_rows()
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["text"], "退租")
        self.assertFalse(rows[1]["departed"])

    def test_household_checkout_includes_active_roommates_but_keeps_old_departure_date(self):
        roommate = self.new_stay("同住人", tenancy=self.tenancy, start_date=date(2026, 7, 1))
        visitor = self.new_stay("未关联合同的暂住人", stay_type="visitor", start_date=None)
        manager = self.new_stay("独立管理员", stay_type="manager", start_date=None)
        old = self.new_stay("已离开同住人", tenancy=self.tenancy, start_date=date(2026, 7, 1),
            end_date=date(2026, 8, 1), is_active=False)
        checkout_tenancy(self.tenancy, checkout_date=date(2026, 10, 1))
        rows = police_report_rows()
        self.assertEqual({r["stay"].pk for r in rows if r["departed"]}, {self.stay.pk, roommate.pk, old.pk, visitor.pk})
        self.assertTrue(all(r["text"] == "退租" for r in rows if r["departed"]))
        manager.refresh_from_db(); self.assertTrue(manager.is_active)
        old.refresh_from_db(); self.assertEqual(old.end_date, date(2026, 8, 1))

    def test_planned_checkout_and_contract_expiry_are_still_current(self):
        set_planned_checkout(self.tenancy, planned_date=date(2026, 10, 3), refund_deposit_amount=0)
        self.assertFalse(police_report_rows()[0]["departed"])

    def test_individual_person_and_room_remove_paths_queue_departure(self):
        for route, args in [("person_stay_remove", [self.person.pk, self.stay.pk]),
                            ("room_person_remove", [self.room.pk, self.stay.pk])]:
            with self.subTest(route=route):
                self.stay.is_active = True; self.stay.save()
                self.assertEqual(self.client.post(reverse(route, args=args)).status_code, 302)
                self.assertEqual(police_report_rows()[0]["text"], "退租")

    def test_renewal_never_reports_administrative_record_as_departure(self):
        renewed = renew_tenancy(self.tenancy, end_date=date(2027, 8, 31), monthly_rent=3000, payment_cycle="monthly")
        rows = police_report_rows()
        self.assertEqual(len(rows), 1); self.assertFalse(rows[0]["departed"])
        self.assertEqual(rows[0]["stay"].tenancy_id, renewed.pk)
        self.stay.refresh_from_db(); self.assertIsNone(self.stay.police_departure_token)

    def test_snapshot_redownload_preserves_generated_content_after_edits_and_confirmation(self):
        self.depart(); report, first = self.generate()
        self.person.name = "已修改姓名"; self.person.save()
        self.confirm(report)
        response = authenticated_client(self).get(reverse("police_report_download", args=[report.pk]))
        original = list(load_workbook(BytesIO(first.content)).active.values)
        restored = list(load_workbook(BytesIO(response.content)).active.values)
        self.assertEqual(original, restored)
        self.assertEqual(PoliceReportExport.objects.count(), 1)
        self.assertContains(self.client.get(reverse("police_report_detail", args=[report.pk])), "测试主租客")
        self.assertEqual(police_report_rows(), [])

    def test_generate_page_redirects_to_its_specific_batch_and_confirmation_requires_post(self):
        self.depart()
        response = self.client.post(reverse("police_preview"))
        report = PoliceReportExport.objects.get()
        self.assertRedirects(response, reverse("police_report_detail", args=[report.pk]) + "?download=1")
        self.assertEqual(self.client.get(reverse("police_report_confirm", args=[report.pk])).status_code, 405)
        self.stay.refresh_from_db(); self.assertIsNone(self.stay.police_departure_reported_at)

    def test_serialization_failure_does_not_create_batch_or_consume_departure(self):
        self.depart()
        with patch("core.exports.Workbook.save", side_effect=OSError("模拟导出失败")):
            with self.assertRaises(OSError): self.client.get(reverse("export_police"))
        self.assertFalse(PoliceReportExport.objects.exists())
        self.assertEqual(police_report_rows()[0]["text"], "退租")

    def test_fingerprint_detects_departure_and_confirmation(self):
        current = police_fingerprint(); self.depart()
        departed = police_fingerprint(); self.assertNotEqual(current, departed)
        report, _ = self.generate(); self.confirm(report)
        self.assertNotEqual(departed, police_fingerprint())

    def test_migration_queues_legacy_departures_once_and_skips_renewals(self):
        renewal = renew_tenancy(self.tenancy, end_date=date(2027, 8, 31), monthly_rent=3000, payment_cycle="monthly")
        old = self.new_stay("旧历史住户", is_active=False, start_date=None, report_note="退租")
        Stay.objects.filter(pk=old.pk).update(police_departure_token=None)
        migration = import_module("core.migrations.0010_police_departure_reporting")
        migration.queue_unconfirmed_departures(apps, None)
        old.refresh_from_db(); token = old.police_departure_token
        self.assertIsNotNone(token)
        migration.queue_unconfirmed_departures(apps, None)
        old.refresh_from_db(); self.assertEqual(old.police_departure_token, token)
        self.stay.refresh_from_db(); self.assertIsNone(self.stay.police_departure_token)
        self.assertEqual(renewal.stays.filter(is_active=True).count(), 1)

    def test_confirmed_inactive_state_is_reported_even_when_old_end_date_is_unknown_or_future(self):
        unknown = self.new_stay("离开日期未知", is_active=False, start_date=None)
        future = self.new_stay("旧记录保留合同截止日", is_active=False, start_date=None, end_date=date(2027, 1, 1))
        rows = {r["stay"].pk: r["text"] for r in police_report_rows()}
        self.assertEqual(rows[unknown.pk], "退租"); self.assertEqual(rows[future.pk], "退租")

    def test_long_note_has_space_to_wrap_and_missing_ids_render_as_blank(self):
        self.new_stay("长备注暂住人", stay_type="visitor", start_date=None, report_note="亲属假期暂住说明" * 20)
        report, response = self.generate()
        sheet = load_workbook(BytesIO(response.content)).active
        long_row = next(row for row in sheet.iter_rows(min_row=3) if row[0].value == "长备注暂住人")
        self.assertGreater(sheet.row_dimensions[long_row[0].row].height, 80)
        self.assertNotContains(self.client.get(reverse("police_report_detail", args=[report.pk])), ">None<")

    def test_clearing_business_data_also_clears_report_snapshots(self):
        self.generate(); clear_business_data()
        self.assertFalse(PoliceReportExport.objects.exists())
        self.generate_empty_report()
        call_command("clear_business_data", yes=True, verbosity=0)
        self.assertFalse(PoliceReportExport.objects.exists())

    def generate_empty_report(self):
        _, response = self.generate()
        self.assertEqual(load_workbook(BytesIO(response.content)).active.max_row, 2)
