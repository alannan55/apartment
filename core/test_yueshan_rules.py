import json
from datetime import date
from importlib import import_module
from io import BytesIO
from unittest.mock import patch

from django.apps import apps
from django.core.exceptions import ValidationError
from django.core.management import call_command, CommandError
from django.test import TestCase
from openpyxl import Workbook

from .exports import police_report_workbook, police_report_rows
from .forms import PersonCreateForm, PersonForm, TenancyEditForm
from .models import Charge, Payment, Person, Room, Stay, Tenancy
from .services import renew_tenancy
from .spreadsheet_import import import_template, TemplateImportError


from .test_support import configure_test_account


class YueshanRuleTests(TestCase):
    def setUp(self):
        configure_test_account(self)
        self.room = Room.objects.create(number="A01", listing_price=3500)
        self.other = Room.objects.create(number="A101", status="self_use")
        self.person = Person.objects.create(name="示例主租客", id_number="110101199001010011")
        self.tenancy = Tenancy.objects.create(
            room=self.room, primary_person=self.person, start_date=date(2026, 6, 1),
            end_date=date(2026, 8, 31), monthly_rent=3000, billing_enabled=False, deposit_amount=None,
        )

    def stay(self, **kwargs):
        fields = dict(person=self.person, room=self.room, tenancy=self.tenancy,
                      start_date=date(2026, 6, 4), end_date=None)
        fields.update(kwargs)
        return Stay.objects.create(**fields)

    def edit_payload(self, **kwargs):
        fields = dict(room=self.room.pk, primary_person=self.person.pk,
                      start_date="2026-06-01", end_date="2026-08-31", monthly_rent="3000",
                      payment_cycle="monthly", status="active", billing_enabled="")
        fields.update(kwargs)
        return fields

    def test_short_main_report_is_one_inclusive_calendar_year(self):
        self.assertEqual(self.tenancy.police_end_date, date(2027, 5, 31))
        self.assertEqual(self.tenancy.contract_text_for_report, "2026.06.01-2027.05.31")
        self.assertEqual(self.tenancy.end_date, date(2026, 8, 31))

    def test_calendar_boundary_and_leap_year(self):
        for start, end, short in [
            (date(2026, 6, 1), date(2026, 11, 29), True),
            (date(2026, 6, 1), date(2026, 11, 30), False),
            (date(2024, 2, 29), date(2024, 8, 27), True),
            (date(2024, 2, 29), date(2024, 8, 28), False),
        ]:
            with self.subTest(start=start, end=end):
                self.tenancy.start_date, self.tenancy.end_date = start, end
                self.assertEqual(self.tenancy.requires_year_police_report, short)
        # Billing's existing rounded-month behavior must stay intact.
        self.tenancy.start_date, self.tenancy.end_date = date(2026, 6, 1), date(2026, 11, 29)
        self.assertFalse(self.tenancy.is_short_term)

    def test_override_is_preserved_but_short_report_is_automatically_one_year(self):
        form = TenancyEditForm(self.edit_payload(police_report_end_date="2026-11-30"), instance=self.tenancy)
        self.assertTrue(form.is_valid(), form.errors)
        form.save()
        self.tenancy.refresh_from_db()
        self.assertEqual(self.tenancy.police_report_end_date, date(2026, 11, 30))
        self.assertEqual(self.tenancy.police_end_date, date(2027, 5, 31))

    def test_edit_rejects_reverse_contract_and_report_dates(self):
        form = TenancyEditForm(self.edit_payload(end_date="2026-05-31"), instance=self.tenancy)
        self.assertFalse(form.is_valid())
        self.tenancy.refresh_from_db()
        form = TenancyEditForm(self.edit_payload(police_report_start_date="2027-06-01", police_report_end_date="2027-05-31"), instance=self.tenancy)
        self.assertFalse(form.is_valid())

    def test_independent_report_start_and_personal_text(self):
        self.tenancy.end_date = date(2026, 11, 30)
        self.tenancy.police_report_start_date = date(2026, 6, 4)
        self.tenancy.police_report_end_date = date(2027, 6, 3)
        self.tenancy.save()
        self.stay(report_note="2026.6.4-2027.6.3租约延长")
        roommate = Person.objects.create(name="示例同住人")
        self.stay(person=roommate, report_note="2026.6.5-2026.8.31同住")
        rows = list(police_report_workbook(date(2026, 10, 1)).active.iter_rows(min_row=3, values_only=True))
        self.assertEqual({r[0]: r[6] for r in rows}, {"示例主租客": "2026.6.4-2027.6.3租约延长", "示例同住人": "2026.6.5-2026.8.31同住"})
        self.assertEqual(self.tenancy.start_date, date(2026, 6, 1))

    def test_editing_earlier_move_in_does_not_detach_current_contract(self):
        self.tenancy.start_date = date(2026, 8, 6)
        self.tenancy.end_date = date(2027, 8, 31)
        self.tenancy.save()
        s = self.stay(start_date=date(2026, 8, 5))
        with patch("django.utils.timezone.localdate", return_value=date(2026, 10, 1)):
            response = self.client.post(f"/people/{self.person.pk}/stays/{s.pk}/edit/", dict(
                room=self.room.pk, stay_type="permanent", start_date="2026-08-05",is_active="on",report_note=""))
        self.assertEqual(response.status_code, 302)
        s.refresh_from_db()
        self.assertEqual(s.tenancy_id, self.tenancy.pk)
        self.assertEqual(s.start_date, date(2026, 8, 5))
        self.assertEqual(s.report_text, self.tenancy.contract_text_for_report)

    def test_historical_stay_edit_keeps_its_original_contract(self):
        self.tenancy.status = "ended"
        self.tenancy.save()
        s = self.stay(is_active=False)
        response = self.client.post(f"/people/{self.person.pk}/stays/{s.pk}/edit/", dict(
            room=self.room.pk, stay_type="permanent", start_date="2026-06-04",report_note=""))
        self.assertEqual(response.status_code,302)
        s.refresh_from_db()
        self.assertEqual(s.tenancy_id,self.tenancy.pk)
        self.assertFalse(s.is_active)

    def test_main_text_override_is_normalized_but_visitor_keeps_temporary_period(self):
        main = self.stay(report_note="2026.6.1-2026.8.31")
        self.assertEqual(main.report_text, "2026.06.01-2027.05.31")
        visitor = Person.objects.create(name="示例探亲人")
        self.stay(person=visitor, tenancy=None, stay_type="visitor", start_date=date(2026, 8, 5), end_date=date(2026, 8, 20), report_note="2026.8.5-2026.8.20母亲探望孩子暂住")
        rows = list(police_report_workbook(date(2026, 10, 1)).active.iter_rows(min_row=3, values_only=True))
        self.assertEqual({r[0]: r[6] for r in rows}[visitor.name], "2026.8.5-2026.8.20母亲探望孩子暂住")

    def test_export_uses_two_queries_and_rejects_invalid_stored_override(self):
        s = self.stay()
        with self.assertNumQueries(2):
            rows = police_report_rows(date(2026, 10, 1))
        self.assertEqual(len(rows), 1)
        # QuerySet updates bypass model save; export must still enforce the rule.
        Stay.objects.filter(pk=s.pk).update(report_note="2026.8.31-2026.6.1")
        with self.assertRaises(ValidationError): police_report_workbook(date(2026, 10, 1))
        with patch("django.utils.timezone.localdate", return_value=date(2026, 10, 1)):
            response = self.client.get("/exports/police.xlsx")
        self.assertRedirects(response, "/people/report/", fetch_redirect_response=False)

    def test_unknown_visitor_dates_are_current_everywhere(self):
        self.person = Person.objects.create(name="示例假期住户")
        s = self.stay(tenancy=None, stay_type="visitor", start_date=None, report_note="假期暂住")
        self.assertTrue(s.active_on(date(2026, 10, 1)))
        with patch("django.utils.timezone.localdate", return_value=date(2026, 10, 1)):
            self.assertContains(self.client.get("/people/"), self.person.name)
            self.assertContains(self.client.get(f"/rooms/{self.room.pk}/"), "日期未提供")
            preview = self.client.get("/people/report/")
            self.assertContains(preview, "假期暂住")
            self.assertEqual(preview.context["report_count"], 1)
        self.assertEqual(police_report_workbook(date(2026, 10, 1)).active.max_row, 3)

    def test_expiry_is_a_reminder_not_departure(self):
        self.person = Person.objects.create(name="示例探望住户")
        s = self.stay(tenancy=None, stay_type="visitor", start_date=date(2026, 8, 5), end_date=date(2026, 8, 20))
        self.assertTrue(s.active_on(date(2026, 10, 1)))
        s.is_active = False; s.save()
        self.assertFalse(s.active_on(date(2026, 10, 1)))

    def test_duplicate_current_identity_is_rejected_without_modifying_person(self):
        self.stay()
        with self.assertRaises(ValidationError): self.stay(room=self.other, tenancy=None)
        response = self.client.post("/people/new/", dict(name="覆盖旧姓名", id_number=self.person.id_number,
                                    room=self.other.pk, stay_type="visitor", report_note="探望"))
        self.assertEqual(response.status_code, 200)
        self.person.refresh_from_db(); self.assertEqual(self.person.name, "示例主租客")
        self.assertEqual(Stay.objects.filter(person=self.person, is_active=True).count(), 1)

    def test_visitor_form_does_not_require_a_fictitious_date(self):
        form = PersonCreateForm(dict(name="示例探望人", id_number="110101199002020022", room=self.room.pk,
                                     stay_type="visitor", report_note="偶尔探望女友"))
        self.assertTrue(form.is_valid(), form.errors)
        self.assertIsNone(form.cleaned_data["start_date"])

    def test_visitor_entry_returns_to_people_without_defaulting_dates(self):
        page = self.client.get("/stays/visitor/new/")
        self.assertIsNone(page.context["form"].initial.get("start_date"))
        response = self.client.post("/stays/visitor/new/", dict(person_name="无日期访客",id_number="110101199004040044",
            room=self.room.pk,report_note="假期暂住",next=page.context["return_url"]))
        self.assertRedirects(response, "/people/", fetch_redirect_response=False)
        self.assertIsNone(Stay.objects.get(person__name="无日期访客").start_date)

    def test_renewal_does_not_copy_stale_personal_report_period(self):
        self.tenancy.end_date = date(2026, 11, 30); self.tenancy.save()
        self.stay(report_note="2026.6.1-2026.11.30")
        with patch("django.utils.timezone.localdate", return_value=date(2026, 12, 1)):
            renewed = renew_tenancy(self.tenancy, end_date=date(2027, 11, 30), monthly_rent=3000, payment_cycle="monthly")
        self.assertEqual(renewed.stays.get(is_active=True).report_note, "")
        self.assertFalse(Charge.objects.exists()); self.assertFalse(Payment.objects.exists())

    def test_template_unknown_visitor_date_and_invalid_period_rollback(self):
        w = Workbook(); w.active.title = "房间与合同"
        w.active.append(["header"]); w.active.append(["A09", "空房"])
        people = w.create_sheet("入住人员"); people.append(["header"])
        people.append(["A09", "暂住", "假期住户", "110101199003030033", "", "", "", "", "", "", "是", "假期暂住"])
        stream=BytesIO(); w.save(stream); stream.seek(0)
        import_template(stream)
        self.assertIsNone(Stay.objects.get(person__name="假期住户").start_date)
        w2=Workbook(); w2.active.title="合同"; w2.active.append(["header"])
        w2.active.append(["A01", self.person.id_number, "2026-07-01", "2026-06-01", "", 3000,"月付",None,None,None,None,"在租","逆序测试"])
        stream=BytesIO(); w2.save(stream); stream.seek(0)
        with self.assertRaises(TemplateImportError): import_template(stream)
        self.assertEqual(Tenancy.objects.count(), 1)

    def test_legacy_import_command_is_disabled_before_writing(self):
        with self.assertRaises(CommandError): call_command("import_current_state", data_dir="missing", clear=True)
        self.assertEqual(Person.objects.count(), 1)

    def test_migration_recovers_explicit_imported_undated_visitor_once(self):
        p=Person.objects.create(name="已导入假期住户", notes='[悦山资料同步 2026-10-01]\n'+json.dumps(
            dict(room="A01", status="active", missing_date=True, residents_period="假期暂住", reporting_period="假期暂住"),ensure_ascii=False))
        migration=import_module("core.migrations.0009_yueshan_rules")
        migration.restore_undated_visitors(apps, None)
        migration.restore_undated_visitors(apps, None)
        s=Stay.objects.get(person=p)
        self.assertIsNone(s.start_date); self.assertEqual(s.report_note,"假期暂住")
        self.assertFalse(Charge.objects.exists()); self.assertFalse(Payment.objects.exists())

    def test_person_notes_are_readable_without_losing_source_metadata(self):
        source = '[悦山资料同步 2026-10-01]\n'+json.dumps(dict(residents_note="拆床板需恢复",room="A01"),ensure_ascii=False)
        self.person.notes = source
        self.person.save()
        self.assertEqual(self.person.display_notes, "拆床板需恢复")
        form = PersonForm(dict(name=self.person.name,id_number=self.person.id_number,notes="已恢复床板"),instance=self.person)
        self.assertTrue(form.is_valid(), form.errors)
        form.save()
        self.person.refresh_from_db()
        self.assertIn(source,self.person.notes)
        self.assertContains(self.client.get(f"/people/{self.person.pk}/"),"已恢复床板")
        self.assertNotContains(self.client.get(f"/people/{self.person.pk}/"),"residents_note")
