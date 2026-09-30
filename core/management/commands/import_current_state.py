import re
from datetime import date
from decimal import Decimal
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from core.models import Charge, Payment, Person, Room, Stay, Tenancy


DATE_RE = re.compile(r"(\d{4})[./-](\d{1,2})[./-](\d{1,2})")


def as_text(value):
    if value is None:
        return ""
    return str(value).strip()


def parse_money(value):
    if value in [None, ""]:
        return None
    text = as_text(value)
    match = re.search(r"\d+(?:\.\d+)?", text)
    return Decimal(match.group(0)) if match else None


def parse_dates(text):
    dates = []
    for year, month, day in DATE_RE.findall(as_text(text)):
        try:
            dates.append(date(int(year), int(month), int(day)))
        except ValueError:
            pass
    return dates


def parse_room_password(text):
    match = re.search(r"密码[:：]\s*([^\s，,]+)", as_text(text))
    return match.group(1) if match else ""


class Command(BaseCommand):
    help = "从现有悦山 Excel 表导入房间、人员和当前入住状态；不导入旧合同和历史账单。"

    def add_arguments(self, parser):
        parser.add_argument("--data-dir", default=r"X:\project\悦山Loft", help="包含现有 Excel 表的目录")
        parser.add_argument("--clear", action="store_true", help="导入前清空业务数据")

    def handle(self, *args, **options):
        data_dir = Path(options["data_dir"])
        files = {
            "rooms": data_dir / "悦山公寓房态表2026.6.xlsx",
            "people": data_dir / "悦山入住人员信息.xlsx",
            "police": data_dir / "给苏警官的6月悦山入住人员信息.xlsx",
        }
        missing = [str(path) for path in files.values() if not path.exists()]
        if missing:
            raise CommandError("找不到文件：\n" + "\n".join(missing))
        with transaction.atomic():
            if options["clear"]:
                Payment.objects.all().delete()
                Charge.objects.all().delete()
                Stay.objects.all().delete()
                Tenancy.objects.all().delete()
                Person.objects.all().delete()
                Room.objects.all().delete()
            rooms = self.import_rooms(files["rooms"])
            people = self.import_people(files["people"])
            stays = self.import_current_stays(files["police"])
            self.stdout.write(self.style.SUCCESS(f"导入完成：房间 {rooms}，人员 {people}，当前入住 {stays}。未导入旧合同和历史账单。"))

    def import_rooms(self, path):
        sheet = load_workbook(path, data_only=True).active
        count = 0
        in_table = False
        for row in sheet.iter_rows(values_only=True):
            if row and as_text(row[0]) == "序号":
                in_table = True
                continue
            if not in_table or not row or not as_text(row[1]):
                continue
            room_number = as_text(row[1])
            status_text = as_text(row[6])
            notes = as_text(row[7])
            if "物业自用" in notes or "自用" in notes:
                status = Room.Status.SELF_USE
            elif "已租" in status_text:
                status = Room.Status.OCCUPIED
            elif "到期" in status_text:
                status = Room.Status.EXPIRING
            else:
                status = Room.Status.VACANT
            Room.objects.update_or_create(
                number=room_number,
                defaults={
                    "orientation": as_text(row[2]),
                    "listing_price": parse_money(row[3]),
                    "floor": as_text(row[4]),
                    "area": parse_money(row[5]),
                    "status": status,
                    "room_password": parse_room_password(notes),
                    "notes": notes,
                },
            )
            count += 1
        return count

    def import_people(self, path):
        sheet = load_workbook(path, data_only=True).active
        count = 0
        for row in sheet.iter_rows(min_row=3, values_only=True):
            name, id_number = as_text(row[0]), as_text(row[1])
            if not name or not id_number:
                continue
            Person.objects.update_or_create(
                id_number=id_number,
                defaults={
                    "name": name,
                    "phone": as_text(row[4]),
                    "emergency_name": as_text(row[5]),
                    "emergency_phone": as_text(row[6]),
                    "notes": as_text(row[7]),
                },
            )
            count += 1
        return count

    def import_current_stays(self, path):
        sheet = load_workbook(path, data_only=True).active
        count = 0
        for row in sheet.iter_rows(min_row=3, values_only=True):
            name, id_number, room_number = as_text(row[0]), as_text(row[1]), as_text(row[2])
            if not name or not id_number or not room_number:
                continue
            person, _ = Person.objects.update_or_create(
                id_number=id_number,
                defaults={
                    "name": name,
                    "phone": as_text(row[3]),
                    "emergency_name": as_text(row[4]),
                    "emergency_phone": as_text(row[5]),
                },
            )
            room = Room.objects.filter(number=room_number).first()
            if not room:
                continue
            contract_text = as_text(row[6])
            stay_type = Stay.Type.VISITOR if ("暂住" in contract_text or "探望" in contract_text) else Stay.Type.PERMANENT
            dates = parse_dates(contract_text)
            start = dates[0] if dates else date.today()
            end = dates[1] if len(dates) > 1 and stay_type == Stay.Type.VISITOR else None
            Stay.objects.update_or_create(
                person=person,
                room=room,
                tenancy=None,
                defaults={
                    "stay_type": stay_type,
                    "start_date": start,
                    "end_date": end,
                    "is_active": True,
                    "report_note": contract_text if stay_type == Stay.Type.VISITOR else "",
                },
            )
            count += 1
        return count
